"""
auto_refresh.py — fully automated HubSpot → Excel refresh.
 
Runs daily in GitHub Actions. Pulls three slices from HubSpot:
  1. NEW SUBMISSIONS      — submission_date within the window
  2. STAGE CHANGES        — hs_v2_date_entered_current_stage within the window
  3. FIELD EDITS          — hs_lastmodifieddate within the window (catches amount,
                            service category, owner, dates changing on previously
                            logged deals without a stage move)
 
Also re-pulls every current Awards-sheet ID so amount / date / service-category
changes on existing awards get in even if hs_lastmodifieddate wasn't bumped.
 
Honors .refresh_work/overrides.json:
  exclude_ids, dealname_prefix, force_remove_from_awards, owner_map
 
Environment:
  HUBSPOT_TOKEN   — HubSpot Private App token (required)
  LOOKBACK_DAYS   — optional, defaults to 2
"""
import json, os, re, sys, time
from datetime import datetime, timedelta, timezone
from pathlib import Path
 
try:
    import openpyxl, requests
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl requests -q")
    import openpyxl, requests
 
# This file is at <repo-root>/.refresh_work/auto_refresh.py
ROOT = Path(__file__).resolve().parent.parent
EXCEL = ROOT / "data" / "CAPS_RFP_Dashboard_Dataset.xlsx"
OVR_PATH = ROOT / ".refresh_work" / "overrides.json"
LAST_PULL_PATH = ROOT / ".refresh_work" / ".last_pull"
 
TOKEN = os.environ.get("HUBSPOT_TOKEN", "").strip()
if not TOKEN:
    print("ERROR: HUBSPOT_TOKEN environment variable not set."); sys.exit(1)
 
LOOKBACK_DAYS = int(os.environ.get("LOOKBACK_DAYS", "2"))
 
STAGE_MAP = {
    "presentationscheduled":"Submitted",
    "1620129473":"Interview",
    "2485737153":"Intent to Award",
    "closedwon":"Closed Won",
    "closedlost":"Closed Lost",
    "2203296493":"Terminated",
    "2766010076":"RFx Cancelled",
    "appointmentscheduled":"Sourced",
    "qualifiedtobuy":"Qualified",
    "1617217262":"Not Qualified",
}
AWARD_STAGES = {"closedwon", "2485737153"}
# Only stage IDs confirmed to exist on the default pipeline. Extra IDs in
# STAGE_MAP are display-only labels for stages that may appear via other paths.
TRACKED_STAGES = [
    "presentationscheduled", "1620129473", "2485737153",
    "closedwon", "closedlost", "2203296493", "2766010076",
]
 
OVR = json.loads(OVR_PATH.read_text()) if OVR_PATH.exists() else {}
EXCLUDE = set(OVR.get("exclude_ids", []))
PREFIX = OVR.get("dealname_prefix", {})
FORCE_RM = set(OVR.get("force_remove_from_awards", []))
OWNER_MAP = OVR.get("owner_map", {})
 
PROPS = [
    "hs_object_id","dealname","agency","submission_date","submission_mode",
    "closedate","createdate","dealstage","hubspot_owner_id","amount",
    "service_category__cloned_","rfp_number","intent_to_awarded_date",
    "awarded_date","interview_date_time","bafo_date","interview_type"
]
 
API = "https://api.hubapi.com"
HEAD = {"Authorization": f"Bearer {TOKEN}", "Content-Type":"application/json"}
 
def to_date(s):
    if not s: return None
    try: return datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
    except: return None
 
def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().lower())
 
def hubspot_search(filter_groups, after=None):
    body = {"filterGroups": filter_groups, "properties": PROPS, "limit": 100}
    if after: body["after"] = after
    r = requests.post(f"{API}/crm/v3/objects/deals/search",
                      headers=HEAD, json=body, timeout=30)
    if r.status_code == 429:
        time.sleep(10); return hubspot_search(filter_groups, after)
    r.raise_for_status()
    return r.json()
 
def search_all(filter_groups):
    out, after = [], None
    while True:
        data = hubspot_search(filter_groups, after)
        out.extend(data.get("results", []))
        paging = data.get("paging", {}).get("next", {})
        if not paging.get("after"): break
        after = paging["after"]
        time.sleep(0.2)
    return out
 
def hubspot_batch_read(ids):
    out = []
    i = 0
    while i < len(ids):
        chunk = [{"id": str(x)} for x in ids[i:i+100]]
        r = requests.post(f"{API}/crm/v3/objects/deals/batch/read",
                          headers=HEAD,
                          json={"properties": PROPS, "inputs": chunk},
                          timeout=30)
        if r.status_code == 429:
            time.sleep(10); continue
        r.raise_for_status()
        out.extend(r.json().get("results", []))
        i += 100
        time.sleep(0.2)
    return out
 
def pull_window(since_iso):
    pipe = {"propertyName":"pipeline","operator":"EQ","value":"default"}
    tracked = {"propertyName":"dealstage","operator":"IN","values":TRACKED_STAGES}
 
    print(f"  Pulling submissions since {since_iso[:10]}...")
    subs = search_all([{"filters":[
        pipe, tracked,
        {"propertyName":"submission_date","operator":"GTE","value":since_iso[:10]},
    ]}])
    print(f"    {len(subs)} submissions")
 
    print(f"  Pulling stage changes since {since_iso}...")
    stages = search_all([{"filters":[
        pipe,
        {"propertyName":"dealstage","operator":"IN","values":[
            "1620129473","2485737153","closedwon","closedlost","2203296493","2766010076"
        ]},
        {"propertyName":"hs_v2_date_entered_current_stage","operator":"GTE","value":since_iso},
    ]}])
    print(f"    {len(stages)} stage changes")
 
    print(f"  Pulling general field edits since {since_iso}...")
    edits = search_all([{"filters":[
        pipe, tracked,
        {"propertyName":"hs_lastmodifieddate","operator":"GTE","value":since_iso},
    ]}])
    print(f"    {len(edits)} modified deals")
 
    merged = {}
    for d in subs + stages + edits:
        did = str(d["id"])
        if did not in merged: merged[did] = d["properties"]
    return merged
 
def refresh_award_rows_from_hubspot(award_ids):
    print(f"  Re-pulling {len(award_ids)} current Award IDs for full refresh...")
    out = hubspot_batch_read(award_ids)
    print(f"    got {len(out)} back")
    return {str(d["id"]): d["properties"] for d in out}
 
def apply_to_excel(fresh):
    wb = openpyxl.load_workbook(EXCEL)
    ws = wb["RFP Data"]
    hdr = [c.value for c in ws[1]]; col = {h:i+1 for i,h in enumerate(hdr)}
 
    xl_row, ag2state, max_sno = {}, {}, 0
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, col["HubSpot ID"]).value
        if rid: xl_row[str(rid).strip()] = r
        ag = ws.cell(r, col["Agency"]).value; st = ws.cell(r, col["Agency State"]).value
        if ag and st: ag2state.setdefault(norm(ag), str(st).strip())
        sn = ws.cell(r, col["S.No."]).value
        if isinstance(sn, (int, float)) and sn > max_sno: max_sno = sn
 
    next_row, next_sno = ws.max_row + 1, max_sno + 1
    updated, inserted, cw_moves = 0, 0, []
 
    for did, p in fresh.items():
        if did in EXCLUDE: continue
        stage_id = p.get("dealstage", "")
        stage = STAGE_MAP.get(stage_id, stage_id)
        owner = OWNER_MAP.get(p.get("hubspot_owner_id", ""), p.get("hubspot_owner_id", ""))
        agency = p.get("agency") or ""
        state = ag2state.get(norm(agency), "")
        amt_raw = p.get("amount")
        amt = float(amt_raw) if amt_raw not in (None, "", "null") else ""
        dn = p.get("dealname") or ""
        if did in PREFIX and not dn.startswith(PREFIX[did]):
            dn = PREFIX[did] + dn
 
        if did in xl_row:
            r = xl_row[did]
            old_stage = ws.cell(r, col["Stage"]).value or ""
            ws.cell(r, col["Stage"]).value = stage
            if dn: ws.cell(r, col["Deal Name"]).value = dn
            if agency: ws.cell(r, col["Agency"]).value = agency
            if state and not (ws.cell(r, col["Agency State"]).value or "").strip():
                ws.cell(r, col["Agency State"]).value = state
            if owner: ws.cell(r, col["Owner"]).value = owner
            if p.get("submission_date"): ws.cell(r, col["Submission Date"]).value = to_date(p["submission_date"])
            if p.get("closedate"): ws.cell(r, col["Bid Closing Date"]).value = to_date(p["closedate"])
            if amt != "": ws.cell(r, col["Amount ($)"]).value = amt
            if p.get("service_category__cloned_"):
                ws.cell(r, col["Service Category"]).value = p["service_category__cloned_"]
            if p.get("submission_mode"): ws.cell(r, col["Submission Mode"]).value = p["submission_mode"]
            if p.get("intent_to_awarded_date"):
                ws.cell(r, col["Intent to Award Date"]).value = to_date(p["intent_to_awarded_date"])
            if p.get("awarded_date"):
                ws.cell(r, col["Awarded Date"]).value = to_date(p["awarded_date"])
            if p.get("interview_date_time"):
                ws.cell(r, col["Interview Date"]).value = to_date(p["interview_date_time"])
            if stage in ("Closed Lost", "Terminated", "RFx Cancelled"):
                ws.cell(r, col["Intent to Award Date"]).value = None
                ws.cell(r, col["Awarded Date"]).value = None
            if stage == "Closed Won" and old_stage != "Closed Won":
                cw_moves.append(did)
            updated += 1
        else:
            rec = {
                "S.No.": next_sno, "HubSpot ID": did, "RFP Number": p.get("rfp_number",""),
                "Deal Name": dn, "Agency": agency, "Agency State": state,
                "Stage": stage, "Interview Flag": "No",
                "Bid Closing Date": to_date(p.get("closedate")),
                "Submission Date": to_date(p.get("submission_date")),
                "Amount ($)": amt, "Service Category": p.get("service_category__cloned_", ""),
                "Submission Mode": p.get("submission_mode", ""), "Owner": owner,
                "Created Date": to_date(p.get("createdate")),
                "HubSpot Link": f"https://app.hubspot.com/contacts/243046792/record/0-3/{did}",
                "Intent to Award Date": to_date(p.get("intent_to_awarded_date")),
                "Awarded Date": to_date(p.get("awarded_date")),
                "Interview Date": to_date(p.get("interview_date_time")),
            }
            for k, v in rec.items():
                if k in col: ws.cell(next_row, col[k]).value = v
            if stage == "Closed Won": cw_moves.append(did)
            next_row += 1; next_sno += 1; inserted += 1
 
    # ----- Awards sheet sync -----
    aws = wb["Awards"]
    ahdr = [c.value for c in aws[1]]; acol = {h:i+1 for i,h in enumerate(ahdr)}
    awards_rows = {}
    for r in range(2, aws.max_row + 1):
        rid = str(aws.cell(r, acol["HubSpot ID"]).value or "").strip()
        if rid: awards_rows[rid] = r
    a_max = max([aws.cell(r, acol["S.No."]).value or 0 for r in range(2, aws.max_row+1)] + [0])
 
    # Remove rows that no longer qualify
    remove_rows = []
    for did, r in awards_rows.items():
        if did in FORCE_RM:
            remove_rows.append((r, did, "force_remove")); continue
        p = fresh.get(did)
        if p and p.get("dealstage") not in AWARD_STAGES:
            remove_rows.append((r, did, f"stage={p.get('dealstage')}"))
    for r, did, why in sorted(remove_rows, key=lambda x: -x[0]):
        print(f"    Removing from Awards: {did} ({why})")
        aws.delete_rows(r, 1); awards_rows.pop(did, None)
 
    # Update existing Award rows
    for did, r in list(awards_rows.items()):
        p = fresh.get(did)
        if not p: continue
        dn = p.get("dealname") or ""
        if did in PREFIX and not dn.startswith(PREFIX[did]):
            dn = PREFIX[did] + dn
        if dn: aws.cell(r, acol["Deal Name"]).value = dn
        if p.get("agency"): aws.cell(r, acol["Agency"]).value = p["agency"]
        aws.cell(r, acol["Stage"]).value = STAGE_MAP.get(p["dealstage"], p["dealstage"])
        amt_raw = p.get("amount")
        if amt_raw not in (None, "", "null"):
            try: aws.cell(r, acol["Amount ($)"]).value = float(amt_raw)
            except: pass
        if p.get("service_category__cloned_") and "Service Category" in acol:
            aws.cell(r, acol["Service Category"]).value = p["service_category__cloned_"]
        own = OWNER_MAP.get(p.get("hubspot_owner_id",""), p.get("hubspot_owner_id",""))
        if own: aws.cell(r, acol["Owner"]).value = own
        if p.get("intent_to_awarded_date"):
            aws.cell(r, acol["Intent to Award Date"]).value = to_date(p["intent_to_awarded_date"])
        if p.get("awarded_date"):
            aws.cell(r, acol["Awarded Date"]).value = to_date(p["awarded_date"])
 
    # Insert any new CW/ITA deal not already in Awards
    for did, p in fresh.items():
        if did in EXCLUDE or did in FORCE_RM: continue
        if p.get("dealstage") not in AWARD_STAGES: continue
        if did in awards_rows: continue
        a_max += 1; r = aws.max_row + 1
        dn = p.get("dealname") or ""
        if did in PREFIX and not dn.startswith(PREFIX[did]):
            dn = PREFIX[did] + dn
        own = OWNER_MAP.get(p.get("hubspot_owner_id",""), p.get("hubspot_owner_id",""))
        agency = p.get("agency") or ""
        state = ag2state.get(norm(agency), "")
        amt_raw = p.get("amount")
        amt = float(amt_raw) if amt_raw not in (None, "", "null") else ""
        rec = {
            "S.No.": a_max, "HubSpot ID": did, "Deal Name": dn, "Agency": agency,
            "Agency State": state, "Stage": STAGE_MAP.get(p["dealstage"], p["dealstage"]),
            "Bid Closing Date": to_date(p.get("closedate")),
            "Submission Date": to_date(p.get("submission_date")),
            "Amount ($)": amt, "Service Category": p.get("service_category__cloned_", ""),
            "Owner": own, "Created Date": to_date(p.get("createdate")),
            "HubSpot Link": f"https://app.hubspot.com/contacts/243046792/record/0-3/{did}",
            "Intent to Award Date": to_date(p.get("intent_to_awarded_date")),
            "Awarded Date": to_date(p.get("awarded_date")),
        }
        for k, v in rec.items():
            if k in acol: aws.cell(r, acol[k]).value = v
        print(f"    Appended to Awards: {did} {agency[:40]}")
 
    # Interview Flag maintenance
    for sheet_name in ("RFP Data", "Awards"):
        s = wb[sheet_name]
        shdr = [c.value for c in s[1]]; scol = {h:i+1 for i,h in enumerate(shdr)}
        if "Interview Flag" not in scol: continue
        for r in range(2, s.max_row + 1):
            iv = s.cell(r, scol.get("Interview Date", 0)).value if "Interview Date" in scol else None
            st = str(s.cell(r, scol.get("Stage", 0)).value or "") if "Stage" in scol else ""
            cur = str(s.cell(r, scol["Interview Flag"]).value or "").strip()
            if (iv or st == "Interview") and cur != "Yes":
                s.cell(r, scol["Interview Flag"]).value = "Yes"
 
    wb.save(EXCEL)
    print(f"\n  Updated {updated}, Inserted {inserted}. New CW moves: {len(cw_moves)}.")
    print(f"  RFP Data rows: {ws.max_row-1}, Awards rows: {aws.max_row-1}")
 
def get_award_ids():
    wb = openpyxl.load_workbook(EXCEL, read_only=True)
    aws = wb["Awards"]
    hdr = [c.value for c in next(aws.iter_rows(min_row=1,max_row=1))]
    cidx = hdr.index("HubSpot ID")
    ids = []
    for row in aws.iter_rows(min_row=2, values_only=True):
        v = row[cidx]
        if v: ids.append(str(v).strip())
    return ids
 
def read_last_pull():
    if LAST_PULL_PATH.exists():
        try: return datetime.fromisoformat(LAST_PULL_PATH.read_text().strip())
        except: pass
    return datetime.now(timezone.utc) - timedelta(days=LOOKBACK_DAYS)
 
def write_last_pull(dt):
    LAST_PULL_PATH.write_text(dt.isoformat())
 
def main():
    now = datetime.now(timezone.utc)
    last = read_last_pull()
    window_start = min(last, now - timedelta(days=LOOKBACK_DAYS))
    print(f"Window: {window_start.isoformat()} -> {now.isoformat()}")
    print(f"Excel : {EXCEL}")
 
    fresh = pull_window(window_start.isoformat())
    aw_ids = get_award_ids()
    aw_fresh = refresh_award_rows_from_hubspot(aw_ids)
    for did, p in aw_fresh.items():
        if did not in fresh: fresh[did] = p
 
    apply_to_excel(fresh)
    write_last_pull(now)
    print(f"\nSaved {EXCEL.name}. Last pull marker updated.")
 
if __name__ == "__main__":
    main()
